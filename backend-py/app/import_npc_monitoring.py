"""Imports real 2026 NPC Budget/IO Budget/IO Actual from the user's NPC
Monitoring workbook (see models_npc_monitoring.py's docstring for why this
exists instead of live workflow tables or the SAP broker). Re-runnable -
deletes the target fiscal year's existing rows first, so pointing this at
next month's file (e.g. "..._Oct'26.xlsx") is a clean refresh.

Run manually from backend-py/, mirroring seed_mock_sap.py's own convention:
    python -m app.import_npc_monitoring "<path to xlsx>" --fiscal-year 2026

The workbook has two parts, read differently:
- Per-SBU tabs (Malls/Estates/Offices/Residential/IT/HR/Admin) list that
  SBU's approved NPC projects - one Budget Code per project, "Revised"
  (Amount minus Cut, confirmed with the user as the right figure) feeding
  NpcMonitoringProject. Column layout differs per tab (confirmed by reading
  every one directly), so columns are located by header text, not a fixed
  index - except HR (an employee car-plan roster, not a per-project table -
  its one pooled NPC budget line is special-cased). Admin used to have no
  Budget Code anywhere in the workbook and was skipped entirely, but as of
  the Sep'26 file it has its own single real project row (Budget Code
  "LPA-2026B-NPC001", confirmed against the Monitoring tab's own
  "CORPORATE - ADMIN" rows too, e.g. AUFNR 40001262 "LAMINATED CABINETS FOR
  MS. CAROL MILLS" carries that same Budget Source Code) - same
  Code/Amount/Title column shape as any other SBU tab, so it now goes
  through the same generic import_sbu_tab path as the rest instead of being
  special-cased or skipped.
- The Monitoring tab lists every real SAP Internal Order - IO code (AUFNR),
  the same Budget Code linking it back to a project, and live
  Budget/Actual/Committed/Allotted/Available - feeding NpcMonitoringIo.
  Uniform layout, fixed columns.
"""

from __future__ import annotations

import argparse
import re
from datetime import datetime

import openpyxl
from sqlmodel import Session, delete

from .db import engine
from .models_npc_monitoring import NpcMonitoringIo, NpcMonitoringProject

# Tab name -> the existing 8-value NPC SBU code (npc_sbu.py). Leisure has no
# tab in the workbook as of Sep'26, but is listed so a future month's file
# adding one needs no script change.
SBU_TAB_TO_NPC_SBU = {
    "Malls": "MALLS",
    "Estates": "ESTATES",
    "Offices": "OFFICES",
    "Residential": "RESIDENTIAL",
    "Leisure": "LEISURE",
    "IT": "CORPORATE_IT",
    "HR": "CORPORATE_HR",
    "Admin": "CORPORATE_ADMIN",
}

# Monitoring tab's own SBU2 column spelling -> the same npc_sbu codes.
MONITORING_SBU2_TO_NPC_SBU = {
    "MALLS": "MALLS",
    "ESTATES": "ESTATES",
    "OFFICES": "OFFICES",
    "RESIDENTIAL": "RESIDENTIAL",
    "LEISURE": "LEISURE",
    "CORPORATE - IT": "CORPORATE_IT",
    "CORPORATE - HR": "CORPORATE_HR",
    "CORPORATE - ADMIN": "CORPORATE_ADMIN",
}

MONITORING_HEADER_ROW = 6
MONITORING_FIRST_DATA_ROW = 7
COL_AUFNR = 1
COL_IO_DESCRIPTION = 2
COL_SBU2 = 4
COL_GROUP = 5
COL_BUDGET_SOURCE_CODE = 8
COL_BUDGET = 26
COL_ACTUAL = 27
COL_COMMITTED = 28
COL_ALLOTTED = 29
COL_AVAILABLE = 30
# The "YTD Forecast" block (columns 48-59, one cumulative-through-that-month
# figure per month, Jan first) - labeled "Forecast" in the workbook itself
# (row 4, confirmed by reading the header cells directly), but its value for
# an already-elapsed month matches this same row's own `actual` figure
# exactly, so it's the source for "IO Actual through month X" (see
# NpcMonitoringIo.ytd_actual_by_month's own docstring).
COL_YTD_FORECAST_JAN = 48

_HAS_LETTER = re.compile(r"[A-Za-z]")


def _clean_title(value: object) -> str:
    """A handful of Malls tab titles start with a leading "*" (e.g.
    "*GH TENANT SPACE RE-CUT & UTILITIES STUB-OUT") - some internal
    flagging convention in the source workbook, not part of the project's
    actual name, so it's stripped for display rather than shown verbatim.
    """
    text = str(value or "").strip()
    return text.lstrip("*").strip()


def _looks_like_a_real_budget_code(value: object) -> bool:
    """Filters out the workbook's own placeholder values in a "Budget
    Code" column - a bare SAP order number (a "Carry-over" row with no
    current-cycle code, e.g. 10001190) or blank. A real code always has at
    least one letter (e.g. "WPD-2026B-NPC001").
    """
    text = str(value or "").strip()
    return bool(text) and bool(_HAS_LETTER.search(text))


def _find_header_row_and_columns(ws, max_scan_rows: int = 10) -> tuple[int, dict[str, int]] | None:
    """Scans the first `max_scan_rows` rows for one containing a cell whose
    text includes "code" (case-insensitive) - that's the header row. Returns
    that row number plus a {lowercased header text: column index} map, or
    None if no such row is found in range.
    """
    for r in range(1, max_scan_rows + 1):
        row_texts = {}
        found_code_col = False
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            text = str(v).strip().lower()
            row_texts[text] = c
            if "code" in text:
                found_code_col = True
        if found_code_col:
            return r, row_texts
    return None


def _find_column(headers: dict[str, int], *needles: str) -> int | None:
    for text, col in headers.items():
        for needle in needles:
            if needle in text:
                return col
    return None


def import_sbu_tab(
    ws, tab_name: str, fiscal_year: int, source_file: str, aufnr_to_budget_code: dict[str, str]
) -> tuple[list[NpcMonitoringProject], dict[str, float]]:
    npc_sbu = SBU_TAB_TO_NPC_SBU.get(tab_name)
    if not npc_sbu:
        print(f"  [{tab_name}] not a recognized SBU tab name - skipped.")
        return [], {}

    header = _find_header_row_and_columns(ws)
    if not header:
        print(f"  [{tab_name}] no header row with a \"Code\" column found in the first 10 rows - skipped.")
        return [], {}
    header_row, headers = header

    code_col = _find_column(headers, "code")
    amount_col = _find_column(headers, "amount")
    cut_col = _find_column(headers, "cut")
    revised_col = _find_column(headers, "revised")
    title_col = _find_column(headers, "project title", "description", "title")
    # Both new as of the Sep'26 file, and only on the Malls tab so far - the
    # rest of the SBU tabs are still on the old layout, so both come back
    # None there and every row's transfer contribution is just 0 (no change
    # in behavior for a tab that hasn't been updated yet).
    io_code_col = _find_column(headers, "io code")
    budget_transfer_col = _find_column(headers, "budget transfer")
    if not code_col or not amount_col:
        print(f"  [{tab_name}] couldn't locate both a Code and an Amount column - skipped.")
        return [], {}

    # Pass 1: one NpcMonitoringProject per anchor row (a real Budget Code
    # with its own Amount) - unchanged from before Budget Transfer existed.
    code_to_project: dict[str, NpcMonitoringProject] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        budget_code_raw = ws.cell(row=r, column=code_col).value
        if not _looks_like_a_real_budget_code(budget_code_raw):
            continue
        budget_code = str(budget_code_raw).strip()
        amount = ws.cell(row=r, column=amount_col).value
        if amount is None:
            continue
        amount = float(amount)
        if revised_col is not None and ws.cell(row=r, column=revised_col).value is not None:
            revised = float(ws.cell(row=r, column=revised_col).value)
        elif cut_col is not None and ws.cell(row=r, column=cut_col).value is not None:
            revised = amount - float(ws.cell(row=r, column=cut_col).value)
        else:
            revised = amount
        title = _clean_title(ws.cell(row=r, column=title_col).value) if title_col else ""
        code_to_project[budget_code] = NpcMonitoringProject(
            fiscal_year=fiscal_year,
            npc_sbu=npc_sbu,
            budget_code=budget_code,
            project_title=title,
            revised_amount=revised,
            source_file=source_file,
        )

    # Pass 1b: a Carry-over project (no current-cycle Budget Code) still
    # gets its own row here, just keyed by its own AUFNR instead of a real
    # code (e.g. "10001072") - _looks_like_a_real_budget_code rejects that
    # as a Budget Code (no letter), which is correct for Pass 1, but its
    # Revised figure is exactly what a Carry-over IO's NPC Budget should
    # show too (see import_monitoring_ios's own docstring / run_import's
    # merge step) - collected here as aufnr->revised, matched up later
    # against the Monitoring tab's own AUFNR for that same IO.
    carry_over_revised: dict[str, float] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        code_raw = ws.cell(row=r, column=code_col).value
        if code_raw is None or _looks_like_a_real_budget_code(code_raw):
            continue
        try:
            aufnr = str(int(code_raw)).zfill(12)
        except (TypeError, ValueError):
            continue
        amount = ws.cell(row=r, column=amount_col).value
        if amount is None:
            continue
        amount = float(amount)
        if revised_col is not None and ws.cell(row=r, column=revised_col).value is not None:
            revised = float(ws.cell(row=r, column=revised_col).value)
        elif cut_col is not None and ws.cell(row=r, column=cut_col).value is not None:
            revised = amount - float(ws.cell(row=r, column=cut_col).value)
        else:
            revised = amount
        carry_over_revised[aufnr] = revised

    # Pass 2: every row carrying a Budget Transfer value adds/deducts from
    # its owning project's NPC Budget - one project can have several such
    # rows (its own anchor row, plus a "continuation" row per IO the
    # transfer moved budget to/from), and a transfer's destination row
    # sometimes leaves this tab's own Code cell blank even though it does
    # belong to a project (confirmed against AUFNR 10001462, transferred out
    # of RFT-2026B-NPC021's own IO 10001387 into a second IO under the same
    # code - the Malls tab's row for it has no Code, but the Monitoring
    # tab's Budget Source Code for that AUFNR says RFT-2026B-NPC021). So the
    # Monitoring tab's own AUFNR->Budget Code mapping is used to attribute
    # each transfer, falling back to this row's own Code cell only when the
    # IO Code is missing or not found there.
    transfer_by_code: dict[str, float] = {}
    unattributed_transfer_total = 0.0
    if budget_transfer_col:
        for r in range(header_row + 1, ws.max_row + 1):
            bt_raw = ws.cell(row=r, column=budget_transfer_col).value
            if bt_raw is None:
                continue
            bt = float(bt_raw)
            if bt == 0:
                continue
            attributed_code = None
            io_code_raw = ws.cell(row=r, column=io_code_col).value if io_code_col else None
            if io_code_raw is not None:
                aufnr = str(int(io_code_raw)).zfill(12)
                attributed_code = aufnr_to_budget_code.get(aufnr)
            if attributed_code is None:
                own_code_raw = ws.cell(row=r, column=code_col).value
                if _looks_like_a_real_budget_code(own_code_raw):
                    attributed_code = str(own_code_raw).strip()
            if attributed_code and attributed_code in code_to_project:
                transfer_by_code[attributed_code] = transfer_by_code.get(attributed_code, 0.0) + bt
            else:
                unattributed_transfer_total += bt

    for budget_code, transfer in transfer_by_code.items():
        code_to_project[budget_code].revised_amount += transfer

    projects = list(code_to_project.values())
    # Transfers reallocate the same approved pool between projects, so the
    # per-project deltas sum to ~0 (an internal reallocation between a
    # project's own IOs nets to exactly 0 and isn't counted here at all) -
    # report the number of projects with a nonzero net change and the total
    # of those net changes' absolute values instead, which stays
    # informative even when the SBU-wide net is zero.
    net_nonzero = {code: v for code, v in transfer_by_code.items() if v != 0}
    transfer_note = f", Budget Transfer changed {len(net_nonzero)} project(s)' NPC Budget by PHP {sum(abs(v) for v in net_nonzero.values()):,.0f} total" if net_nonzero else ""
    if unattributed_transfer_total:
        transfer_note += f" (PHP {unattributed_transfer_total:,.0f} in Budget Transfer couldn't be attributed to any project - left out)"
    print(f"  [{tab_name}] {len(projects)} project(s) imported{transfer_note}.")
    return projects, carry_over_revised


def import_hr_pooled_project(wb, fiscal_year: int, source_file: str) -> list[NpcMonitoringProject]:
    """HR's tab is an employee car-plan roster, not a per-project table -
    all of it funds one pooled NPC budget line. Amount comes from the tab's
    own row-1 total (the same "Car Plan Cost" column total the tab already
    shows); the budget code is the one seen throughout the Monitoring tab's
    "Corporate - HR" rows.
    """
    if "HR" not in wb.sheetnames:
        return []
    ws = wb["HR"]
    # HR's own header row (row 2) has no "code" column at all, so the
    # generic _find_header_row_and_columns heuristic (which requires one)
    # doesn't apply here - read row 2 directly instead.
    headers = {str(ws.cell(row=2, column=c).value or "").strip().lower(): c for c in range(1, ws.max_column + 1)}
    car_plan_cost_col = _find_column(headers, "car plan cost")
    total = ws.cell(row=1, column=car_plan_cost_col).value if car_plan_cost_col else None
    if total is None:
        print("  [HR] couldn't find a Car Plan Cost total on row 1 - skipped.")
        return []
    project = NpcMonitoringProject(
        fiscal_year=fiscal_year,
        npc_sbu="CORPORATE_HR",
        budget_code="MGM-2026B-NPC001",
        project_title="Company Car Plan (pooled)",
        revised_amount=float(total),
        source_file=source_file,
    )
    print(f"  [HR] 1 pooled project imported (PHP {total:,.0f}).")
    return [project]


def import_monitoring_ios(ws, fiscal_year: int, source_file: str) -> list[NpcMonitoringIo]:
    ios: list[NpcMonitoringIo] = []
    carry_over_count = 0
    skipped_unmapped_sbu = 0
    for r in range(MONITORING_FIRST_DATA_ROW, ws.max_row + 1):
        aufnr_raw = ws.cell(row=r, column=COL_AUFNR).value
        if aufnr_raw is None:
            continue
        sbu2_raw = str(ws.cell(row=r, column=COL_SBU2).value or "").strip().upper()
        npc_sbu = MONITORING_SBU2_TO_NPC_SBU.get(sbu2_raw)
        if not npc_sbu:
            skipped_unmapped_sbu += 1
            continue
        # A blank Budget Source Code means this IO is a "Carry-over" from a
        # prior cycle with no current-year Budget Code - imported anyway
        # (npc_utilization surfaces these as their own standalone rows at
        # the end of the table, per the user's request), just with no
        # project to group it under.
        budget_code_raw = ws.cell(row=r, column=COL_BUDGET_SOURCE_CODE).value
        budget_code = str(budget_code_raw).strip() if budget_code_raw and str(budget_code_raw).strip() else None
        if budget_code is None:
            carry_over_count += 1
        aufnr = str(int(aufnr_raw)).zfill(12)
        ytd_actual_by_month = {
            str(m): float(ws.cell(row=r, column=COL_YTD_FORECAST_JAN + m - 1).value or 0)
            for m in range(1, 13)
        }
        ios.append(
            NpcMonitoringIo(
                fiscal_year=fiscal_year,
                npc_sbu=npc_sbu,
                aufnr=aufnr,
                budget_code=budget_code,
                io_description=str(ws.cell(row=r, column=COL_IO_DESCRIPTION).value or "").strip(),
                budget=float(ws.cell(row=r, column=COL_BUDGET).value or 0),
                actual=float(ws.cell(row=r, column=COL_ACTUAL).value or 0),
                committed=float(ws.cell(row=r, column=COL_COMMITTED).value or 0),
                allotted=float(ws.cell(row=r, column=COL_ALLOTTED).value or 0),
                available=float(ws.cell(row=r, column=COL_AVAILABLE).value or 0),
                ytd_actual_by_month=ytd_actual_by_month,
                group_name=str(ws.cell(row=r, column=COL_GROUP).value or "").strip() or None,
                source_file=source_file,
            )
        )
    print(f"  [Monitoring] {len(ios)} IO(s) imported ({carry_over_count} of them Carry-over, no Budget Code), {skipped_unmapped_sbu} skipped (unrecognized SBU2).")
    return ios


def run_import(wb, fiscal_year: int, source_file: str) -> dict:
    """Does the actual import against an already-open workbook - shared by
    both the CLI entry point below and the Admin Console's upload endpoint
    (routers/npc_monitoring.py), so "override the file initially uploaded"
    from the UI runs the exact same logic as running this script by hand.
    Deletes and replaces every row for `fiscal_year`, then returns a plain
    summary dict (JSON-serializable) rather than relying on the caller to
    scrape printed output.
    """
    # Read before the SBU tabs, not after - Budget Transfer attribution
    # (import_sbu_tab) needs this AUFNR->Budget Code mapping to figure out
    # which project a transfer row belongs to, since a transfer's
    # destination row doesn't always repeat its own tab's Code cell.
    print("Monitoring tab:")
    if "Monitoring" not in wb.sheetnames:
        raise ValueError("No 'Monitoring' tab found in this workbook.")
    ios = import_monitoring_ios(wb["Monitoring"], fiscal_year, source_file)
    aufnr_to_budget_code = {io.aufnr: io.budget_code for io in ios if io.budget_code}

    projects: list[NpcMonitoringProject] = []
    carry_over_revised: dict[str, float] = {}
    print("Per-SBU tabs:")
    for tab_name in ["Malls", "Estates", "Offices", "Residential", "Leisure", "IT", "Admin"]:
        if tab_name in wb.sheetnames:
            tab_projects, tab_carry_over_revised = import_sbu_tab(wb[tab_name], tab_name, fiscal_year, source_file, aufnr_to_budget_code)
            projects.extend(tab_projects)
            carry_over_revised.update(tab_carry_over_revised)
    projects.extend(import_hr_pooled_project(wb, fiscal_year, source_file))

    # An IO with its own numeric-code "Carryover" row in an SBU tab carries
    # the same Revised (Amount - Cut) figure every other project's NPC
    # Budget comes from - preferred here over the Monitoring tab's own
    # `budget` (a live SAP release figure that can genuinely differ, e.g.
    # AUFNR 10001294: SBU tab Revised PHP 3,660,714.29 vs Monitoring tab
    # Budget PHP 4,560,714.29 for the same project) so a Carry-over row's
    # NPC Budget is computed the same way as every other project's instead
    # of substituting a live SAP figure. Deliberately NOT gated on
    # `io.budget_code is None` - AUFNR 10001296 still has its own PHP
    # 48,435.32 "Carryover" line in the Malls tab even though the
    # Monitoring tab now also attributes it to JLC-2026B-NPC003, and
    # matching the SBU tab's own total exactly (per the user's explicit
    # request) means counting it here too, same as the SBU tab itself
    # still does. Falls back to `budget` (unchanged) when no SBU tab has a
    # matching row for that AUFNR.
    for io in ios:
        if io.aufnr in carry_over_revised:
            io.carry_over_revised_amount = carry_over_revised[io.aufnr]

    # Computed before the session closes and expires every attribute on
    # commit - these are plain values already known from construction, not
    # anything that needs a fresh DB read.
    by_sbu: dict[str, int] = {}
    for p in projects:
        by_sbu[p.npc_sbu] = by_sbu.get(p.npc_sbu, 0) + 1

    with Session(engine) as session:
        session.exec(delete(NpcMonitoringProject).where(NpcMonitoringProject.fiscal_year == fiscal_year))
        session.exec(delete(NpcMonitoringIo).where(NpcMonitoringIo.fiscal_year == fiscal_year))
        for p in projects:
            session.add(p)
        for io in ios:
            session.add(io)
        session.commit()

    print(f"\nDone. {len(projects)} project(s), {len(ios)} IO(s) imported for fiscal year {fiscal_year}.")
    print("Projects by SBU:", by_sbu)

    return {
        "sourceFile": source_file,
        "fiscalYear": fiscal_year,
        "projectCount": len(projects),
        "ioCount": len(ios),
        "projectsBySbu": by_sbu,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", help="Path to the NPC Monitoring .xlsx file")
    parser.add_argument("--fiscal-year", type=int, default=2026, help="Fiscal year these figures belong to (default 2026)")
    args = parser.parse_args()

    source_file = args.path.split("\\")[-1].split("/")[-1]
    print(f"Reading {source_file} for fiscal year {args.fiscal_year}...")
    wb = openpyxl.load_workbook(args.path, read_only=True, data_only=True)
    run_import(wb, args.fiscal_year, source_file)


if __name__ == "__main__":
    main()
